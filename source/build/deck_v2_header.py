# -*- coding: utf-8 -*-
"""WeYield — Building the AI-Native Business. McKinsey-grammar deck, python-pptx, shape-drawn exhibits."""
import re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

OUT = sys.argv[1] if len(sys.argv) > 1 else "WeYield - Building the AI-Native Business.pptx"

# ---------- palette ----------
NAVY = RGBColor(0x05, 0x1C, 0x2C)
BLUE = RGBColor(0x22, 0x51, 0xFF)
LBLUE = RGBColor(0x99, 0xB4, 0xFF)
PALE = RGBColor(0xE8, 0xEE, 0xFF)
INK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x7F, 0x7F, 0x7F)
MID = RGBColor(0xB3, 0xB3, 0xB3)
RULE = RGBColor(0xD9, 0xD9, 0xD9)
BG2 = RGBColor(0xF4, 0xF5, 0xF7)
WARN = RGBColor(0xB3, 0x66, 0x1A)
WARN_SOFT = RGBColor(0xF7, 0xE9, 0xDA)
OK = RGBColor(0x1D, 0x70, 0x49)
BAD = RGBColor(0xA8, 0x34, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB5, 0x1D, 0x1D)
AMBER = RGBColor(0xB3, 0x6B, 0x00)
GREEN = RGBColor(0x1E, 0x7F, 0x4F)

SERIF = "Georgia"
SANS = "Arial"

W, H = 13.333, 7.5
ML, MR = 0.5, 0.5
TITLE_Y, TITLE_H, TITLE_W = 0.42, 1.0, 10.6
BODY_Y, BODY_B = 1.62, 6.75
SRC_Y = 6.95

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]
page_no = [0]


# ---------- primitives ----------
def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(line_w)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = 0.12
        except Exception:
            pass
    s.text_frame.text = ""
    return s


def hline(slide, x, y, w, color=RULE, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln


def vline(slide, x, y, h, color=RULE, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln


def _runs(p, text, size, color, bold=False, italic=False, font=SANS):
    """**bold** and *italic* mini-markup."""
    parts = re.split(r'(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*))', text)
    for t in parts:
        if not t:
            continue
        b, it = bold, italic
        if t.startswith('**') and t.endswith('**'):
            t = t[2:-2]; b = True
        elif t.startswith('*') and t.endswith('*') and len(t) > 2:
            t = t[1:-1]; it = True
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.bold = b; r.font.italic = it
        r.font.color.rgb = color; r.font.name = font


def set_bullet(p, char="•", indent_in=0.18, color=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', str(int(Inches(indent_in))))
    pPr.set('indent', str(-int(Inches(indent_in))))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buClr', 'a:buFont'):
        e = pPr.find(qn(tag))
        if e is not None:
            pPr.remove(e)
    if color is not None:
        buClr = etree.SubElement(pPr, qn('a:buClr'))
        etree.SubElement(buClr, qn('a:srgbClr'), val='%02X%02X%02X' % (color[0], color[1], color[2]))
    buFont = etree.SubElement(pPr, qn('a:buFont'), typeface=SANS)
    etree.SubElement(pPr, qn('a:buChar'), char=char)


def text(slide, x, y, w, h, paras, size=11, color=INK, bold=False, italic=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4, line_spacing=None,
         bullets=False, bullet_char="•", margins=(0, 0, 0, 0), fill=None, line=None):
    """paras: str or list of str | (str, dict overrides)."""
    if fill is not None or line is not None:
        box = rect(slide, x, y, w, h, fill=fill, line=line)
        tf = box.text_frame
    else:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for item in paras:
        if isinstance(item, tuple):
            t, o = item
        else:
            t, o = item, {}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = o.get('align', align)
        p.space_after = Pt(o.get('space_after', space_after))
        if o.get('space_before') is not None:
            p.space_before = Pt(o['space_before'])
        ls = o.get('line_spacing', line_spacing)
        if ls:
            p.line_spacing = ls
        _runs(p, t, o.get('size', size), o.get('color', color), o.get('bold', bold), o.get('italic', italic), o.get('font', font))
        if o.get('bullet', bullets):
            set_bullet(p, o.get('bullet_char', bullet_char), o.get('indent', 0.18), o.get('bullet_color'))
    return box


def cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Each side: None → no line; else (hex, pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        e = tcPr.find(qn(tag))
        if e is not None:
            tcPr.remove(e)
    order = [('a:lnL', left), ('a:lnR', right), ('a:lnT', top), ('a:lnB', bottom)]
    idx = 0
    for tag, spec in order:
        if spec is None:
            ln = etree.Element(qn(tag), w='0')
            etree.SubElement(ln, qn('a:noFill'))
        else:
            hexc, pt = spec
            ln = etree.Element(qn(tag), w=str(int(pt * 12700)), cap='flat', cmpd='sng', algn='ctr')
            sf = etree.SubElement(ln, qn('a:solidFill'))
            etree.SubElement(sf, qn('a:srgbClr'), val=hexc)
            etree.SubElement(ln, qn('a:prstDash'), val='solid')
        tcPr.insert(idx, ln); idx += 1


def table(slide, x, y, w, rows, col_w, size=9.5, header=True, row_h=0.36, header_fill=NAVY,
          header_color=WHITE, zebra=False, bold_first_col=False, valign=MSO_ANCHOR.TOP, pad=0.06,
          col_align=None, header_size=None, status_col=None, header_h=0.42):
    nrows, ncols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r in range(nrows):
        tbl.rows[r].height = Inches(header_h if (header and r == 0) else row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(pad)
            cell.margin_top = Inches(0.045); cell.margin_bottom = Inches(0.045)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE if (header and r == 0) else valign
            val = rows[r][c]
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            if col_align and col_align[c]:
                p.alignment = col_align[c]
            is_hdr = header and r == 0
            if is_hdr:
                cell.fill.solid(); cell.fill.fore_color.rgb = header_fill
                _runs(p, str(val), header_size or size, header_color, bold=True)
                cell_borders(cell, bottom=('051C2C', 0.75))
            else:
                if zebra and r % 2 == 0:
                    cell.fill.solid(); cell.fill.fore_color.rgb = BG2
                else:
                    cell.fill.background()
                colr = INK
                b = bold_first_col and c == 0
                if status_col is not None and c == status_col:
                    v = str(val).upper()
                    if v.startswith('HELD') or v.startswith('COVERED'):
                        colr = GREEN
                    elif 'MISSING' in v or v.startswith('NOT'):
                        colr = RED
                    elif 'ONLY' in v:
                        colr = AMBER
                    elif 'UNKNOWN' in v:
                        colr = BLUE
                    else:
                        colr = INK
                    b = True
                _runs(p, str(val), size, colr, bold=b)
                cell_borders(cell, bottom=('D9D9D9', 0.75))
    return shp


# ---------- chrome ----------
def new_slide(title, tracker=None, sticker=None, source=None):
    s = prs.slides.add_slide(BLANK)
    page_no[0] += 1
    text(s, ML, TITLE_Y, TITLE_W, TITLE_H, title, size=20, color=NAVY, font=SERIF, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.05, space_after=0)
    hline(s, ML, 1.47, W - ML - MR, color=NAVY, weight=0.75)
    if tracker:
        text(s, 9.9, 0.2, 2.93, 0.25, tracker.upper(), size=7.5, color=GREY, align=PP_ALIGN.RIGHT)
    if sticker:
        b = rect(s, 11.53, 0.52, 1.3, 0.27, fill=None, line=NAVY, line_w=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = b.text_frame; tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        _runs(p, sticker.upper(), 7.5, NAVY, bold=True)
    if source:
        text(s, ML, SRC_Y, 10.8, 0.4, source, size=7.5, color=GREY, line_spacing=1.0, space_after=0)
    text(s, 12.0, SRC_Y + 0.02, 0.83, 0.25, str(page_no[0]), size=8.5, color=GREY, align=PP_ALIGN.RIGHT)
    hline(s, ML, SRC_Y - 0.08, W - ML - MR, color=RULE, weight=0.5)
    return s


def callout(slide, x, y, w, h, title, body, fill=PALE, accent=BLUE, title_size=10, body_size=10):
    rect(slide, x, y, w, h, fill=fill)
    rect(slide, x, y, 0.06, h, fill=accent)
    paras = []
    if title:
        paras.append((title, {'bold': True, 'size': title_size, 'color': NAVY, 'space_after': 5}))
    if isinstance(body, str):
        body = [body]
    for b in body:
        paras.append((b, {'size': body_size, 'color': INK, 'space_after': 5}))
    text(slide, x + 0.2, y + 0.14, w - 0.34, h - 0.24, paras, line_spacing=1.1)


def kpi(slide, x, y, w, h, value, label, sub=None, accent=NAVY):
    rect(slide, x, y, w, h, fill=WHITE, line=RULE)
    rect(slide, x, y, w, 0.05, fill=accent)
    text(slide, x + 0.15, y + 0.2, w - 0.3, 0.5, value, size=22, color=NAVY, font=SERIF, space_after=0)
    paras = [(label, {'size': 9, 'color': INK, 'bold': True, 'space_after': 1})]
    if sub:
        paras.append((sub, {'size': 8, 'color': GREY}))
    text(slide, x + 0.15, y + 0.78, w - 0.3, h - 0.85, paras, line_spacing=1.05)


def hbars(slide, x, y, w, items, max_val, label_w=2.6, bar_h=0.34, gap=0.2, fmt=lambda v: f"{v}",
          value_w=1.1, label_size=10, threshold=None, threshold_label=None, label_color=INK):
    """items: (label, value, color, optional value-label override)."""
    bar_x = x + label_w + 0.1
    bar_w_max = w - label_w - 0.1 - value_w
    for i, it in enumerate(items):
        label, val, col = it[0], it[1], it[2]
        vlabel = it[3] if len(it) > 3 else fmt(val)
        yy = y + i * (bar_h + gap)
        text(slide, x, yy, label_w, bar_h, label, size=label_size, color=label_color, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        bw = max(0.02, bar_w_max * val / max_val)
        rect(slide, bar_x, yy, bw, bar_h, fill=col)
        text(slide, bar_x + bw + 0.08, yy, value_w + 0.5, bar_h, vlabel, size=label_size, color=NAVY, bold=True,
             anchor=MSO_ANCHOR.MIDDLE, space_after=0)
    total_h = len(items) * (bar_h + gap) - gap
    vline(slide, bar_x, y - 0.05, total_h + 0.1, color=NAVY, weight=0.75)
    if threshold is not None:
        tx = bar_x + bar_w_max * threshold / max_val
        ln = slide.shapes.add_connector(1, Inches(tx), Inches(y - 0.12), Inches(tx), Inches(y + total_h + 0.12))
        ln.line.color.rgb = RED; ln.line.width = Pt(1.25); ln.line.dash_style = 4
        if threshold_label:
            text(slide, tx - 1.0, y + total_h + 0.16, 2.0, 0.3, threshold_label, size=8.5, color=RED, align=PP_ALIGN.CENTER, bold=True)
    return bar_x, bar_w_max


def vcols(slide, x, y, w, h, items, max_val, fmt=lambda v: f"{v}", col_w=None, top_labels=None,
          label_size=10, value_size=12, base_label=None):
    """items: (label, value, color). Draws baseline + columns + value labels above + category labels below."""
    n = len(items)
    slot = w / n
    cw = col_w or slot * 0.52
    base_y = y + h - 0.5
    plot_h = h - 0.5 - 0.45
    hline(slide, x, base_y, w, color=NAVY, weight=0.75)
    for i, (label, val, col) in enumerate(items):
        cx = x + slot * i + (slot - cw) / 2
        ch = plot_h * val / max_val
        rect(slide, cx, base_y - ch, cw, ch, fill=col)
        text(slide, cx - 0.3, base_y - ch - 0.36, cw + 0.6, 0.32, fmt(val), size=value_size, color=NAVY, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, space_after=0)
        if top_labels and top_labels[i]:
            text(slide, cx - 0.3, base_y - ch - 0.66, cw + 0.6, 0.3, top_labels[i], size=9, color=BLUE, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, space_after=0)
        text(slide, x + slot * i, base_y + 0.06, slot, 0.44, label, size=label_size, color=INK, align=PP_ALIGN.CENTER,
             line_spacing=1.0, space_after=0)


# =====================================================================
# SLIDES
# =====================================================================

